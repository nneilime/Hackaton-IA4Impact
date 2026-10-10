"""Export of the canonical store in the LPC Dataverse target structure (schema/dataverse_tables.xlsx, "le modèle qui fait foi").

One CSV per Dataverse table, columns named with the logical names of the workbook (publisher prefix lpc_), choice
columns carrying the CSV values of the `Choices` sheet, alternate keys filled. Traceability fields and validation
statuses are never dropped. Where our model is richer than the target (related products, evidence excerpts,
extraction method, EPPO status), the extra columns are APPENDED with the suffix `_ext` so an import can ignore them.

`validate_against_workbook()` reads the workbook itself and checks our export against it: every target column present,
business-required columns filled, choice values legal, text lengths respected.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import re
import zipfile
from pathlib import Path

from sqlalchemy.orm import Session

from .config import DISCLAIMER, ROOT
from .db.models import AuthorizedUse, ChangeLog, Product, ProductAlias, Registration, Source
from .normalize import canonical_active_entity

WORKBOOK = ROOT / "schema" / "dataverse_tables.xlsx"
LALLEMAND = re.compile(r"DANSTAR|LALLEMAND", re.I)

# our alias relations -> Dataverse choice lpc_namerelationship (exact value kept in lpc_relationshipdetail_ext)
NAME_RELATION = {"alternate_name": "alternate_name", "historical_name": "historical_or_market_name", "related_formulation": "related_formulation",
                 "sister_product_candidate": "sister_product_candidate", "confirmed_sister_product": "confirmed_sister_product",
                 "related_product": "related_formulation"}


def _d(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "value"):
        return v.value
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%dT%H:%M:%SZ") if (v.hour or v.minute or v.second or v.tzinfo) else v.strftime("%Y-%m-%d")
    return str(v)


def _date(v) -> str:
    return v.strftime("%Y-%m-%d") if v else ""


def _code(prefix: str, raw: str, limit: int = 50) -> str:
    """Stable primary-name code within the Dataverse length limit: raw id when it fits, else prefix + sha1 of the id."""
    if len(raw) <= limit:
        return raw
    return f"{prefix}-{hashlib.sha1(raw.encode('utf-8')).hexdigest()[:16]}"


def _owner(p: Product) -> str:
    holders = " ".join((r.holder_original or "") for r in p.registrations)
    return "lpc" if LALLEMAND.search(holders) or p.product_id.startswith("P-LPC-") or LALLEMAND.search(p.commercial_name) else "competitor"


def _confidence_status(a: ProductAlias) -> str:
    ev = (a.evidence or "").lower()
    if "lpc answer" in ev or "lpc starter kit" in ev and a.review_status.value == "validated":
        return "confirmed_internal"
    if a.relationship_type.value in ("alternate_name", "historical_name") and ("ppid" in ev or "ppls" in ev or "e-phy" in ev or "agrofit" in ev):
        return "confirmed_related"
    return "strong_candidate" if (a.confidence or 0) >= 0.8 else "candidate"


def _evidence_type(a: ProductAlias) -> str:
    ev = (a.evidence or "").lower()
    if "lpc answer" in ev:
        return "lpc_written_statement"
    if "starter kit" in ev:
        return "internal_product_list"
    if "same active entity" in ev:
        return "data_inference_species_strain"
    return "official_record_field"


def build_tables(session: Session) -> dict[str, list[dict]]:
    T: dict[str, list[dict]] = {k: [] for k in ("lpc_product", "lpc_productname", "lpc_activeentity", "lpc_source", "lpc_registration",
                                               "lpc_authorizeduse", "lpc_eppocode", "lpc_changelog")}
    products = session.query(Product).order_by(Product.product_id).all()
    for p in products:
        T["lpc_product"].append({"lpc_productcode": p.product_id, "lpc_commercialname": p.commercial_name, "lpc_formulationcode": p.formulation_code or "",
                                 "lpc_owner": _owner(p), "lpc_issynthetic": _d(p.is_synthetic), "lpc_validationstatus": p.validation_status.value,
                                 "lpc_activeentitynormalized_ext": p.active_entity_normalized or "", "lpc_strain_ext": p.strain or ""})
        T["lpc_productname"].append({"lpc_mappingcode": _code("PN", f"PN-{p.product_id}-CANON"), "lpc_productid": p.product_id, "lpc_namevariant": p.commercial_name,
                                     "lpc_relationshiptype": "canonical_name", "lpc_market": "multiple" if len({r.country_code for r in p.registrations}) > 1
                                     else (next(iter({r.country_code for r in p.registrations}), "global")), "lpc_confidencescore": "1.00",
                                     "lpc_confidencestatus": "confirmed_related", "lpc_evidencesummary": "Primary commercial name as written in the official source(s)",
                                     "lpc_evidencetype": "official_record_field", "lpc_reviewstatus": "automatically_extracted", "lpc_relationshipdetail_ext": "canonical_name",
                                     "lpc_sourcerecordurl_ext": "", "lpc_aliasid_ext": f"PN-{p.product_id}-CANON"})
        for a in sorted(p.aliases, key=lambda x: x.alias_id):
            T["lpc_productname"].append({"lpc_mappingcode": _code("PN", a.alias_id), "lpc_productid": p.product_id, "lpc_namevariant": a.alias_name,
                                         "lpc_relationshiptype": NAME_RELATION[a.relationship_type.value], "lpc_market": a.country_code or "multiple",
                                         "lpc_confidencescore": f"{a.confidence:.2f}" if a.confidence is not None else "", "lpc_confidencestatus": _confidence_status(a),
                                         "lpc_evidencesummary": (a.evidence or "")[:2000], "lpc_evidencetype": _evidence_type(a), "lpc_reviewstatus": a.review_status.value,
                                         "lpc_relationshipdetail_ext": a.relationship_type.value, "lpc_sourcerecordurl_ext": a.source_record_url or "", "lpc_aliasid_ext": a.alias_id})
        # active entities: one row per distinct original wording seen for this product
        seen: set[str] = set()
        originals = [(p.active_entity_original, None)] + [(r.active_entity_original, r) for r in p.registrations]
        n = 0
        for original, reg in originals:
            if not original or original in seen:
                continue
            seen.add(original)
            n += 1
            sp, st, norm, note = canonical_active_entity(original)
            synonym = note is not None
            T["lpc_activeentity"].append({"lpc_mappingcode": _code("AE", f"AE-{p.product_id}-{n}"), "lpc_productid": p.product_id,
                                          "lpc_entitytype": "taxonomic_synonym" if synonym else "microorganism", "lpc_namevariant": original[:300],
                                          "lpc_canonicalspecies": sp or "", "lpc_strain": st or "", "lpc_eppocode": "",
                                          "lpc_relationshiptype": "former_taxonomic_name" if synonym else "current_label_name",
                                          "lpc_confidencescore": "0.90" if synonym else "1.00", "lpc_confidencestatus": "strong_candidate" if synonym else "confirmed_related",
                                          "lpc_evidencesummary": (f"{note}; " if note else "") + (f"as written in {reg.source_id} record {reg.registration_number}" if reg else "as written in the source"),
                                          "lpc_reviewstatus": "to_verify" if synonym else "automatically_extracted"})
    for s in session.query(Source).order_by(Source.source_id):
        T["lpc_source"].append({"lpc_sourcecode": s.source_id, "lpc_officialname": s.official_name, "lpc_countryorregion": s.country_or_region,
                                "lpc_baseurl": s.base_url or "", "lpc_accessmethod": s.access_method.replace("/", ";"), "lpc_dataformat": (s.data_format or "")[:200],
                                "lpc_licensename": (s.license_name or "")[:200], "lpc_accesslimits": (s.access_and_reuse_limits or "")[:2000],
                                "lpc_manualvalidationrequired": _d(s.manual_validation_required)})
    for r in session.query(Registration).order_by(Registration.registration_id):
        T["lpc_registration"].append({"lpc_registrationcode": r.registration_id, "lpc_productid": r.product_id, "lpc_countrycode": r.country_code,
                                      "lpc_authorityname": r.authority_name[:200], "lpc_registrationnumber": r.registration_number or "",
                                      "lpc_statusoriginal": (r.status_original or "")[:200], "lpc_status": r.status_normalized.value,
                                      "lpc_approvaldate": _date(r.approval_date), "lpc_expirationdate": _date(r.expiration_date), "lpc_sourceid": r.source_id,
                                      "lpc_sourcerecordurl": r.source_record_url, "lpc_collectedat": _d(r.collected_at), "lpc_validationstatus": r.validation_status.value,
                                      "lpc_issynthetic": _d(r.is_synthetic),
                                      "lpc_registrationholder_ext": r.holder_original or "", "lpc_activeentityoriginal_ext": r.active_entity_original or "",
                                      "lpc_formulationoriginal_ext": r.formulation_original or "", "lpc_extractionmethod_ext": r.extraction_method.value,
                                      "lpc_evidenceexcerpt_ext": (r.evidence_excerpt or "")[:4000], "lpc_snapshotid_ext": str(r.snapshot_id), "lpc_iscurrent_ext": _d(r.is_current)})
        for u in sorted(r.uses, key=lambda x: x.use_id):
            T["lpc_authorizeduse"].append({"lpc_usecode": u.use_id, "lpc_registrationid": r.registration_id, "lpc_croporiginal": (u.crop_name_original or "")[:300],
                                           "lpc_cropnormalized": (u.crop_name_normalized or "")[:200], "lpc_eppocropcode": u.eppo_crop_code or "",
                                           "lpc_targetoriginal": (u.target_name_original or "")[:300], "lpc_targetnormalized": (u.target_name_normalized or "")[:200],
                                           "lpc_eppotargetcode": u.eppo_target_code or "", "lpc_applicationmethod": (u.application_method or "")[:200],
                                           "lpc_dosevalue": _d(u.dose_value), "lpc_doseunitoriginal": (u.dose_unit_original or "")[:50],
                                           "lpc_doseunitnormalized": u.dose_unit_normalized or "", "lpc_maxapplications": _d(u.max_applications),
                                           "lpc_intervaldays": _d(u.application_interval_days), "lpc_phidays": _d(u.phi_days), "lpc_reihours": _d(u.rei_hours),
                                           "lpc_restrictionsoriginal": (u.restrictions_original or "")[:4000], "lpc_validationstatus": u.validation_status.value,
                                           "lpc_issynthetic": _d(u.is_synthetic),
                                           "lpc_dosevaluenormalized_ext": _d(u.dose_value_normalized), "lpc_eppocropstatus_ext": u.eppo_crop_status or "",
                                           "lpc_eppotargetstatus_ext": u.eppo_target_status or "", "lpc_pairing_ext": u.pairing, "lpc_extractionmethod_ext": u.extraction_method.value,
                                           "lpc_confidence_ext": _d(u.confidence), "lpc_evidenceexcerpt_ext": (u.evidence_excerpt or "")[:4000],
                                           "lpc_sourcerecordurl_ext": u.source_record_url, "lpc_collectedat_ext": _d(u.collected_at)})
    for kind, fn in (("crop", "eppo_crop_mapping.csv"), ("pest", "eppo_pest_mapping.csv")):
        with (ROOT / "data" / fn).open(encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                T["lpc_eppocode"].append({"lpc_conceptcode": row["concept_id"], "lpc_concepttype": kind, "lpc_language": row.get("language", ""),
                                          "lpc_nameoriginal": row["name_original"], "lpc_namenormalized": row["name_normalized"],
                                          "lpc_scientificname": row.get("scientific_name", ""), "lpc_eppocode": row.get("eppo_code", ""),
                                          "lpc_mappingstatus": row.get("mapping_status", "to_verify"), "lpc_multilingualaliases": row.get("multilingual_aliases", "")})
    obs = ROOT / "data" / "eppo_observed_aliases.csv"
    if obs.exists():
        known = {(r["lpc_conceptcode"], r["lpc_concepttype"]) for r in T["lpc_eppocode"]}
        extra: dict[tuple, dict] = {}
        with obs.open(encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                key = (row["concept_id"], row["kind"])
                if key in known:
                    for r in T["lpc_eppocode"]:
                        if (r["lpc_conceptcode"], r["lpc_concepttype"]) == key and row["alias"] not in r["lpc_multilingualaliases"]:
                            r["lpc_multilingualaliases"] = (r["lpc_multilingualaliases"] + "|" if r["lpc_multilingualaliases"] else "") + f"{row['language']}:{row['alias']} (seen in {row['seen_in_source']})"
                else:
                    e = extra.setdefault(key, {"lpc_conceptcode": row["concept_id"], "lpc_concepttype": row["kind"], "lpc_language": row["language"],
                                               "lpc_nameoriginal": row["alias"], "lpc_namenormalized": row["concept_id"].replace("_", " ").capitalize(),
                                               "lpc_scientificname": row.get("scientific_name", ""), "lpc_eppocode": row.get("eppo_code_candidate", ""),
                                               "lpc_mappingstatus": "to_verify", "lpc_multilingualaliases": ""})
                    e["lpc_multilingualaliases"] = (e["lpc_multilingualaliases"] + "|" if e["lpc_multilingualaliases"] else "") + f"{row['language']}:{row['alias']} (seen in {row['seen_in_source']})"
        T["lpc_eppocode"] += list(extra.values())
    for c in session.query(ChangeLog).order_by(ChangeLog.change_id):
        reg = session.get(Registration, c.entity_id) if c.entity_type == "registration" else None
        T["lpc_changelog"].append({"lpc_changecode": f"CHG-{c.change_id:06d}", "lpc_registrationid": c.entity_id if c.entity_type == "registration" else "",
                                   "lpc_fieldname": c.field, "lpc_oldvalue": (c.old_value or "")[:4000], "lpc_newvalue": (c.new_value or "")[:4000],
                                   "lpc_detectedat": _d(c.detected_at), "lpc_sourceid": reg.source_id if reg else ""})
    return T


# --------------------------------------------------------------------------- workbook-driven validation
def load_workbook_model() -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    cols: dict[str, list[dict]] = {}
    for row in wb["Columns"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        cols.setdefault(row[0], []).append({"name": row[1], "type": row[3] or "", "fmt": str(row[4] or ""), "req": row[5] or "", "csv": row[6] or ""})
    choices: dict[str, set[str]] = {}
    for row in wb["Choices"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            choices.setdefault(row[0], set()).add(str(row[2]))
    keys = {row[0]: str(row[2]) for row in wb["Keys"].iter_rows(min_row=2, values_only=True) if row[0]}
    return {"columns": cols, "choices": choices, "keys": keys}


def validate_against_workbook(tables: dict[str, list[dict]]) -> dict:
    model = load_workbook_model()
    report = {"tables": {}, "errors": []}
    for table, cols in model["columns"].items():
        rows = tables.get(table, [])
        info = {"rows": len(rows), "missing_columns": [], "required_empty": {}, "bad_choices": {}, "too_long": {}}
        for c in cols:
            name = c["name"]
            if rows and name not in rows[0]:
                info["missing_columns"].append(name)
                continue
            if c["req"] == "Business required":
                empty = sum(1 for r in rows if r.get(name, "") == "")
                if empty and not (table == "lpc_changelog" and name == "lpc_changecode"):
                    info["required_empty"][name] = empty
            if c["type"] == "Choice" and c["fmt"] in model["choices"]:
                allowed = model["choices"][c["fmt"]]
                bad = sorted({r[name] for r in rows if r.get(name) and not set(r[name].split(";")) <= allowed})
                if bad:
                    info["bad_choices"][name] = bad
            if c["type"].startswith("Single line") and c["fmt"].isdigit():
                lim = int(c["fmt"])
                long = sum(1 for r in rows if len(r.get(name, "")) > lim)
                if long:
                    info["too_long"][name] = long
        if info["missing_columns"] or info["required_empty"] or info["bad_choices"] or info["too_long"]:
            report["errors"].append(table)
        report["tables"][table] = info
    report["ok"] = not report["errors"]
    return report


# --------------------------------------------------------------------------- writers
def write_csvs(tables: dict[str, list[dict]], out_dir: Path) -> dict[str, int]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    counts = {}
    for name, rows in tables.items():
        cols = list(rows[0].keys()) if rows else [c["name"] for c in load_workbook_model()["columns"].get(name, [])]
        with (out_dir / f"{name}.csv").open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)
        counts[name] = len(rows)
    (out_dir / "README.txt").write_text(f"Dataverse-shaped export generated {dt.datetime.now(dt.timezone.utc):%Y-%m-%d %H:%M} UTC.\n"
                                        f"Logical names follow schema/dataverse_tables.xlsx; *_ext columns are additive extensions.\n{DISCLAIMER}\n", encoding="utf-8")
    return counts


def zip_bytes(tables: dict[str, list[dict]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, rows in tables.items():
            s = io.StringIO()
            cols = list(rows[0].keys()) if rows else [c["name"] for c in load_workbook_model()["columns"].get(name, [])]
            w = csv.DictWriter(s, fieldnames=cols, lineterminator="\n")
            w.writeheader()
            w.writerows(rows)
            z.writestr(f"{name}.csv", "﻿" + s.getvalue())
        z.writestr("README.txt", f"Dataverse-shaped export. Logical names follow schema/dataverse_tables.xlsx; *_ext columns are additive.\n{DISCLAIMER}\n")
    return buf.getvalue()


def xlsx_bytes(tables: dict[str, list[dict]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for name, rows in tables.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        cols = list(rows[0].keys()) if rows else [c["name"] for c in load_workbook_model()["columns"].get(name, [])]
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
    ws = wb.create_sheet("README")
    ws.append(["Dataverse-shaped export; logical names from schema/dataverse_tables.xlsx; *_ext columns are additive extensions."])
    ws.append([DISCLAIMER])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
