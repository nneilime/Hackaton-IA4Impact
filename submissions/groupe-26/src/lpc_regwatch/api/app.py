"""FastAPI implementation of schema/openapi.yaml (+ documented extensions).

Run: uvicorn lpc_regwatch.api.app:app --reload   (or scripts/serve_api.py)
Every regulatory answer carries evidence (source id, URL, collection time, validation status) and the disclaimer.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import logging
from typing import Iterator

from datetime import datetime, timedelta, timezone
from fastapi import Depends, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response, StreamingResponse, HTMLResponse
from sqlalchemy.orm import Session

from .. import __version__
from ..config import ROOT, DISCLAIMER, settings
from ..db import get_session
from ..db.models import AuthorizedUse, Product, ProductAlias, Registration, ReviewLog, Snapshot, ValidationStatus
from . import queries as Q
from .ask import ask as run_ask
from .schemas import (AliasOut, Answer, AskRequest, ChangeOut, ComparisonResponse, ProductsResponse, Registration as RegistrationOut,
                      RegistrationsResponse, ReviewRequest)

log = logging.getLogger(__name__)

DESCRIPTION = """Regulatory-watch API for plant protection products (Lallemand Plant Care challenge).

**This API is not an official regulatory source.** Every answer carries `evidence[]` (source id, official URL,
collection time UTC, validation status) and a `disclaimer`. `warnings[]` lists missing or uncertain data.
A 404 or an empty list is **never** proof of no registration (the product may be absent from the collected
source, registered under another name, or the market may not be collected at all).

Sources: Health Canada PPID (CA), US EPA PPLS (US), ANSES E-Phy (FR), MAPA AGROFIT (BR).
"""

from contextlib import asynccontextmanager


@asynccontextmanager
async def lifespan(_: FastAPI):
    """Create tables and seed the sources table on first start (idempotent, additive); the API then serves an
    empty but valid store until the pipeline has run."""
    from ..db import init_db

    try:
        init_db()
    except Exception as e:  # noqa: BLE001 - the health endpoint will show the problem
        log.error("init_db failed at startup: %s", e)
    yield


# Behind the reverse proxy the API lives under /api (the proxy strips the prefix); root_path makes Swagger UI fetch
# /api/openapi.json instead of /openapi.json (which the proxy would route to the UI).
app = FastAPI(title="LPC Global Regulatory AI API", version=__version__, description=DESCRIPTION, root_path=settings.api_root_path, root_path_in_servers=False,
              servers=[{"url": settings.public_base_url}], docs_url="/docs", redoc_url="/redoc", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


def db() -> Iterator[Session]:
    s = get_session()
    try:
        yield s
    finally:
        s.close()


def _countries(csv_: str) -> list[str]:
    out = [c.strip().upper() for c in csv_.split(",") if c.strip()]
    bad = [c for c in out if len(c) != 2 or not c.isalpha()]
    if bad or not out:
        raise HTTPException(422, f"countries must be ISO alpha-2 codes separated by commas, got {bad or csv_!r}")
    return list(dict.fromkeys(out))


# --------------------------------------------------------------------------- contract endpoints
GUIDE_CSS = ("body{font-family:Segoe UI,Arial,sans-serif;color:#212529;background:#f5f6f7;margin:0}main{max-width:900px;margin:0 auto;padding:32px 24px 60px;background:#fff}"
             "h1{font-size:1.7rem;margin-top:0}h2{font-size:1.2rem;margin-top:2rem;border-top:1px solid #e6e8ea;padding-top:1rem}table{border-collapse:collapse;width:100%;font-size:.92rem}"
             "th,td{border:1px solid #e1e3e4;padding:6px 8px;text-align:left;vertical-align:top}th{background:#f1f3f4}code{background:#f1f3f4;padding:1px 4px;border-radius:4px}"
             "blockquote{border-left:4px solid #e30327;margin:1rem 0;padding:.4rem 1rem;color:#5b6770;background:#fafafa}a{color:#36434b}")


@app.get("/guide", include_in_schema=False, response_class=HTMLResponse)
def guide():
    """Human-readable integration guide (docs/API_GUIDE.md): how to use the data in M365, Dataverse, Excel or any program."""
    import markdown

    md = (ROOT / "docs" / "API_GUIDE.md").read_text(encoding="utf-8")
    body = markdown.markdown(md, extensions=["tables", "fenced_code"])
    return HTMLResponse(f"<!doctype html><html lang='fr'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
                        f"<title>Regulatory Watch — guide d'intégration</title><style>{GUIDE_CSS}</style></head><body><main>{body}</main></body></html>")


@app.get("/health", operation_id="getHealth", tags=["contract"])
def health(s: Session = Depends(db)):
    n = s.query(Registration).count()
    last = s.query(Snapshot).filter_by(status="ok").order_by(Snapshot.collected_at.desc()).first()
    return {"status": "ok", "version": __version__, "registrations": n, "aiEnabled": settings.ai_enabled, "aiProvider": settings.ai_provider if settings.ai_enabled else None,
            "aiModel": settings.ai_model if settings.ai_enabled else None, "lastSnapshotAt": last.collected_at if last else None,
            "updateSchedule": schedule_info(last.collected_at if last else None), "frEphyScope": settings.fr_ephy_scope, "disclaimer": DISCLAIMER}


def schedule_info(last: datetime | None) -> dict:
    """The nightly pipeline schedule (cron in UTC, from PIPELINE_SCHEDULE) with the next run, so the UI can say when data is refreshed."""
    cron = settings.pipeline_schedule.strip()
    parts = cron.split()
    nxt = None
    freq = "custom"
    try:
        minute, hour = int(parts[0]), int(parts[1])
        if parts[2:5] == ["*", "*", "*"]:
            freq = "daily"
            now = datetime.now(timezone.utc)
            nxt = now.replace(hour=hour, minute=minute, second=0, microsecond=0)
            if nxt <= now:
                nxt += timedelta(days=1)
        elif parts[2:4] == ["*", "*"] and parts[4] != "*":
            freq = "weekly"
    except (ValueError, IndexError):
        pass
    return {"cron": cron, "timezone": "UTC", "frequency": freq, "nextRunAt": nxt, "lastRunAt": last}


@app.get("/products", operation_id="searchProducts", response_model=ProductsResponse, tags=["contract"])
def products(q: str | None = Query(None, description="Commercial name, alias (official or hypothesis) or active entity; substring, case-insensitive"),
             limit: int = Query(50, ge=1, le=500), s: Session = Depends(db)):
    items = Q.search_products(s, q, limit)
    warnings = [] if items else [f"No product matching '{q}' in the collected sources. Not proof of no registration: try an alias, the active entity, or check sources coverage (/sources)."]
    return {"items": items, "warnings": warnings, "disclaimer": DISCLAIMER}


@app.get("/products/{productId}/registrations", operation_id="getProductRegistrations", response_model=RegistrationsResponse, tags=["contract"],
         responses={404: {"description": "No match. This is not proof of no registration."}})
def product_registrations(productId: str, country: str | None = Query(None, pattern="^[A-Za-z]{2}$"), s: Session = Depends(db)):
    p = Q.resolve_product(s, productId)
    if p is None:
        raise HTTPException(404, {"message": f"Unknown product '{productId}'. This is not proof of no registration.", "disclaimer": DISCLAIMER})
    items = Q.registrations_for(s, p.product_id, country)
    warnings = Q.quality_warnings(items)
    if country and not items:
        _, cov_w = Q.coverage_for(s, p.product_id, [country.upper()], [])
        warnings += cov_w
    return {"items": items, "warnings": warnings, "disclaimer": DISCLAIMER}


@app.get("/compare", operation_id="compareProductUses", response_model=ComparisonResponse, tags=["contract"])
def compare(productId: str, countries: str = Query(..., examples=["CA,US,FR,BR"]), maxUsesPerRegistration: int | None = Query(None, ge=1, le=5000),
            s: Session = Depends(db)):
    p = Q.resolve_product(s, productId)
    if p is None:
        raise HTTPException(404, {"message": f"Unknown product '{productId}'. This is not proof of no registration.", "disclaimer": DISCLAIMER})
    return Q.compare(s, p.product_id, _countries(countries), maxUsesPerRegistration)


@app.post("/ask", operation_id="askRegulatoryQuestion", response_model=Answer, tags=["contract"])
def ask(req: AskRequest, s: Session = Depends(db)):
    return run_ask(s, req.question, req.language, req.context, req.history)


# --------------------------------------------------------------------------- extensions (documented in m365/README.md)
@app.get("/products/{productId}/aliases", response_model=list[AliasOut], tags=["extensions"],
         summary="Official alternate/historical names and sister-product hypotheses with review status")
def aliases(productId: str, relationship: str | None = None, s: Session = Depends(db)):
    p = Q.resolve_product(s, productId)
    if p is None:
        raise HTTPException(404, {"message": f"Unknown product '{productId}'.", "disclaimer": DISCLAIMER})
    return Q.aliases_for(s, p.product_id, relationship)


@app.get("/products/{productId}/related", tags=["extensions"], summary="Distinct products linked to this one (sister, related) with their registrations")
def related(productId: str, countries: str | None = None, s: Session = Depends(db)):
    p = Q.resolve_product(s, productId)
    if p is None:
        raise HTTPException(404, {"message": f"Unknown product '{productId}'.", "disclaimer": DISCLAIMER})
    return {"productId": p.product_id, "items": Q.related_products(s, p.product_id, _countries(countries) if countries else None), "disclaimer": DISCLAIMER}


@app.get("/products/{productId}/quality", tags=["extensions"], summary="Missing or uncertain fields per registration")
def quality(productId: str, s: Session = Depends(db)):
    p = Q.resolve_product(s, productId)
    if p is None:
        raise HTTPException(404, {"message": f"Unknown product '{productId}'.", "disclaimer": DISCLAIMER})
    return Q.data_quality(s, p.product_id)


@app.get("/active-entities/{species}/registrations", tags=["extensions"], summary="Products containing a micro-organism, optionally on a crop")
def by_active_entity(species: str, crop: str | None = None, countries: str | None = None, s: Session = Depends(db)):
    rows = Q.products_by_active_entity(s, species, crop, _countries(countries) if countries else None)
    w = [] if rows else [f"No product with '{species}' in the collected sources - not proof of absence."]
    return {"items": rows, "warnings": w, "disclaimer": DISCLAIMER}


@app.get("/expiring", tags=["extensions"], summary="Registrations expiring within N months (CA and FR carry dates)")
def expiring_(months: int = Query(12, ge=1, le=120), s: Session = Depends(db)):
    rows = Q.expiring(s, months)
    return {"items": rows, "warnings": ["Only CA-PPID and FR-EPHY publish expiry/withdrawal dates; US and BR registrations cannot appear here."],
            "disclaimer": DISCLAIMER}


@app.get("/changes", response_model=list[ChangeOut], tags=["extensions"], summary="Field-level changes detected between snapshots")
def changes_(limit: int = Query(200, ge=1, le=2000), s: Session = Depends(db)):
    return Q.changes(s, limit)


@app.get("/overview", tags=["extensions"], summary="Map-ready overview per country: counts by status, registrations in scope, coverage")
def overview_(scope: str = Query("pilots", pattern="^(pilots|lallemand|all|product)$"), productId: str | None = None,
              months: int = Query(12, ge=1, le=120), s: Session = Depends(db)):
    pid = None
    if scope == "product":
        p = Q.resolve_product(s, productId or "")
        if p is None:
            raise HTTPException(404, {"message": f"Unknown product '{productId}'. This is not proof of no registration.", "disclaimer": DISCLAIMER})
        pid = p.product_id
    return Q.overview(s, scope, pid, months)


@app.get("/sources", tags=["extensions"], summary="Registered official sources, licences, limits and last collection")
def sources(s: Session = Depends(db)):
    return {"items": Q.sources_summary(s), "disclaimer": DISCLAIMER}


@app.post("/review", tags=["extensions"], summary="Regulatory Affairs review: change a validation status (logged)")
def review(req: ReviewRequest, s: Session = Depends(db)):
    model = {"registration": Registration, "use": AuthorizedUse, "alias": ProductAlias, "product": Product}[req.entityType]
    obj = s.get(model, req.entityId)
    if obj is None:
        raise HTTPException(404, f"{req.entityType} {req.entityId} not found")
    field = "review_status" if req.entityType == "alias" else "validation_status"
    old = getattr(obj, field).value
    setattr(obj, field, ValidationStatus(req.newStatus))
    s.add(ReviewLog(entity_type=req.entityType, entity_id=req.entityId, old_status=old, new_status=req.newStatus, reviewer=req.reviewer, note=req.note))
    s.commit()
    return {"entityType": req.entityType, "entityId": req.entityId, "oldStatus": old, "newStatus": req.newStatus, "reviewer": req.reviewer,
            "reviewedAt": dt.datetime.now(dt.timezone.utc), "disclaimer": DISCLAIMER}


@app.get("/review/log", tags=["extensions"], summary="Human review history")
def review_log(limit: int = Query(200, ge=1, le=2000), s: Session = Depends(db)):
    return [{"reviewId": r.review_id, "entityType": r.entity_type, "entityId": r.entity_id, "oldStatus": r.old_status, "newStatus": r.new_status,
             "reviewer": r.reviewer, "note": r.note, "reviewedAt": r.reviewed_at} for r in s.query(ReviewLog).order_by(ReviewLog.review_id.desc()).limit(limit)]


EXPORT_COLUMNS = ["product_id", "commercial_name", "country_code", "registration_number", "status_original", "status_normalized",
                  "approval_date", "expiration_date", "holder_original", "active_entity_original", "crop_name_original", "crop_name_normalized",
                  "eppo_crop_code", "eppo_crop_status", "target_name_original", "target_name_normalized", "eppo_target_code", "application_method",
                  "dose_value", "dose_unit_original", "dose_value_normalized", "dose_unit_normalized", "max_applications", "phi_days",
                  "restrictions_original", "pairing", "extraction_method", "validation_status", "source_id", "source_record_url", "collected_at"]


def _export_rows(s: Session, product_id: str | None, countries: list[str] | None) -> Iterator[list]:
    q = s.query(Registration).join(Product)
    if product_id:
        q = q.filter(Registration.product_id == product_id)
    if countries:
        q = q.filter(Registration.country_code.in_(countries))
    for r in q.order_by(Registration.product_id, Registration.country_code):
        base = [r.product_id, r.product.commercial_name, r.country_code, r.registration_number, r.status_original, r.status_normalized.value,
                r.approval_date.strftime("%Y-%m-%d") if r.approval_date else "", r.expiration_date.strftime("%Y-%m-%d") if r.expiration_date else "",
                r.holder_original, r.active_entity_original]
        uses = r.uses or [None]
        for u in uses:
            if u is None:
                row = base + [""] * 16
            else:
                row = base + [u.crop_name_original, u.crop_name_normalized, u.eppo_crop_code, u.eppo_crop_status, u.target_name_original,
                              u.target_name_normalized, u.eppo_target_code, u.application_method, u.dose_value, u.dose_unit_original,
                              u.dose_value_normalized, u.dose_unit_normalized, u.max_applications, u.phi_days, u.restrictions_original, u.pairing]
            row += [(u.extraction_method.value if u else r.extraction_method.value), (u.validation_status.value if u else r.validation_status.value),
                    r.source_id, r.source_record_url, r.collected_at.strftime("%Y-%m-%dT%H:%M:%SZ")]
            yield ["" if v is None else v for v in row]


@app.get("/export.csv", tags=["extensions"], summary="Flat product x country x use table (for Power Automate / SharePoint / Excel)")
def export_csv_(productId: str | None = None, countries: str | None = None, s: Session = Depends(db)):
    pid = None
    if productId:
        p = Q.resolve_product(s, productId)
        if p is None:
            raise HTTPException(404, {"message": f"Unknown product '{productId}'.", "disclaimer": DISCLAIMER})
        pid = p.product_id
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(EXPORT_COLUMNS)
    for row in _export_rows(s, pid, _countries(countries) if countries else None):
        w.writerow(row)
    data = buf.getvalue().encode("utf-8-sig")
    return Response(content=data, media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="lpc_regwatch_{pid or "all"}.csv"', "X-Disclaimer": DISCLAIMER})


@app.get("/export.xlsx", tags=["extensions"], summary="Same table as Excel, with a Sources sheet and the disclaimer")
def export_xlsx(productId: str | None = None, countries: str | None = None, s: Session = Depends(db)):
    from openpyxl import Workbook

    pid = None
    if productId:
        p = Q.resolve_product(s, productId)
        if p is None:
            raise HTTPException(404, {"message": f"Unknown product '{productId}'.", "disclaimer": DISCLAIMER})
        pid = p.product_id
    wb = Workbook()
    ws = wb.active
    ws.title = "comparison"
    ws.append(EXPORT_COLUMNS)
    for row in _export_rows(s, pid, _countries(countries) if countries else None):
        ws.append(row)
    ws2 = wb.create_sheet("sources")
    ws2.append(["source_id", "official_name", "country", "base_url", "license", "limits", "registrations", "last_collected_at"])
    for src in Q.sources_summary(s):
        ws2.append([src["sourceId"], src["officialName"], src["countryOrRegion"], src["baseUrl"], src["licenseName"], src["accessAndReuseLimits"],
                    src["registrations"], src["lastCollectedAt"].isoformat() if src["lastCollectedAt"] else ""])
    ws3 = wb.create_sheet("disclaimer")
    ws3.append([DISCLAIMER])
    ws3.append([f"Generated {dt.datetime.now(dt.timezone.utc):%Y-%m-%d %H:%M} UTC by LPC Regulatory Watch v{__version__}"])
    out = io.BytesIO()
    wb.save(out)
    return StreamingResponse(iter([out.getvalue()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="lpc_regwatch_{pid or "all"}.xlsx"'})


# --------------------------------------------------------------------------- Dataverse target structure (schema/dataverse_tables.xlsx)
@app.get("/export/dataverse.zip", tags=["m365"], summary="All tables in the LPC Dataverse target structure (one CSV per table, logical names lpc_*)")
def export_dataverse_zip(s: Session = Depends(db)):
    from ..dataverse import build_tables, zip_bytes

    data = zip_bytes(build_tables(s))
    return Response(content=data, media_type="application/zip", headers={"Content-Disposition": 'attachment; filename="lpc_dataverse_export.zip"'})


@app.get("/export/dataverse.xlsx", tags=["m365"], summary="Same tables as one workbook (one sheet per Dataverse table)")
def export_dataverse_xlsx(s: Session = Depends(db)):
    from ..dataverse import build_tables, xlsx_bytes

    return Response(content=xlsx_bytes(build_tables(s)), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="lpc_dataverse_export.xlsx"'})


@app.get("/export/dataverse/{table}", tags=["m365"], summary="One Dataverse table as JSON rows (lpc_product, lpc_registration, lpc_authorizeduse, ...)")
def export_dataverse_table(table: str, limit: int = Query(1000, ge=1, le=20000), offset: int = Query(0, ge=0), s: Session = Depends(db)):
    from ..dataverse import build_tables

    tables = build_tables(s)
    if table not in tables:
        raise HTTPException(404, {"message": f"Unknown table '{table}'. Tables: {', '.join(tables)}", "disclaimer": DISCLAIMER})
    rows = tables[table]
    return {"table": table, "total": len(rows), "offset": offset, "items": rows[offset:offset + limit], "disclaimer": DISCLAIMER}


@app.get("/export/dataverse-validation", tags=["m365"], summary="Check the export against schema/dataverse_tables.xlsx (columns, required, choices, lengths)")
def export_dataverse_validation(s: Session = Depends(db)):
    from ..dataverse import build_tables, validate_against_workbook

    return validate_against_workbook(build_tables(s))


# --------------------------------------------------------------------------- LPC internal registry reconciliation
@app.get("/registry/sample", tags=["registry"], summary="Synthetic internal-registry rows (is_synthetic=true) in LPC's SharePoint column layout")
def registry_sample():
    from ..registry import normalize_registry_rows, sample_rows

    return {"columns": list(sample_rows()[0].keys()), "rows": normalize_registry_rows(sample_rows()), "disclaimer": DISCLAIMER}


@app.get("/registry/reconcile/sample", tags=["registry"], summary="Reconcile the synthetic registry sample with the official sources")
def registry_reconcile_sample(s: Session = Depends(db)):
    from ..registry import normalize_registry_rows, reconcile, sample_rows

    return reconcile(s, normalize_registry_rows(sample_rows()))


@app.post("/registry/reconcile", tags=["registry"], summary="Upload a CSV/XLSX export of the LPC registry list and reconcile it with the official sources")
async def registry_reconcile(file: UploadFile = File(...), s: Session = Depends(db)):
    from ..registry import normalize_registry_rows, read_registry_file, reconcile

    data = await file.read()
    if len(data) > 20_000_000:
        raise HTTPException(413, "file too large (20 MB max)")
    try:
        rows = read_registry_file(data, file.filename or "registry.csv")
    except Exception as e:  # noqa: BLE001
        raise HTTPException(422, f"cannot read file: {type(e).__name__}") from e
    missing = [c for c in ("Title", "Country", "Registration Status", "Registration Number") if rows and c not in rows[0]]
    if missing:
        raise HTTPException(422, f"missing registry columns: {missing}")
    return reconcile(s, normalize_registry_rows(rows))


@app.exception_handler(HTTPException)
async def http_exc(_, exc: HTTPException):
    detail = exc.detail if isinstance(exc.detail, dict) else {"message": str(exc.detail), "disclaimer": DISCLAIMER}
    return JSONResponse(status_code=exc.status_code, content=detail)
