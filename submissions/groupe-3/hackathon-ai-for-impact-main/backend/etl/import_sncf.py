"""Import des VRAIS fichiers SNCF du hackathon (data/, jamais commité) dans SQLite.

Fichiers (dans data/ à la racine du repo, ou backend/data/, ou $LIO_DATA_DIR) :
- Data_Hackathon.csv            indicateurs mensuels par train (« ; », BOM, pourcentages « 63,20% ») ;
- « Détail desserte (1).xlsx »  circulations et arrêts ordonnés, février à juin 2026 (lecture en flux) ;
- annuaire_etablissement.xlsx   établissements scolaires (table `etablissement`, Occitanie) ;
- reseau_urbain_gare_coordonnees/stops.txt (GTFS, à la racine du repo) : coordonnées des gares,
  rapprochées par nom normalisé ; à défaut, coordonnées du réseau de démo (etl/reseau_demo.py).

La lecture du xlsx (800 000 lignes) prend environ 40 s : la base SQLite construite sert de cache.

HYPOTHÈSE « SEMAINE TYPE » : la desserte couvre février–juin 2026 mais le scénario de démonstration
est le 12 novembre 2026. Pour chaque jour de FENETRE_SEMAINE_TYPE, les circulations sont recopiées
(est_simule = 1) depuis le jour de la même position dans la semaine le plus récent présent dans la
desserte (jour_type_pour). `hypotheses()` la décrit et /api/sante l'expose.

Usage (depuis backend/) : python -m etl.import_sncf
"""
from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path

from app import config, db
from . import commun, import_datatourisme, reseau_demo

FICHIER_DESSERTE = "Détail desserte (1).xlsx"
FICHIER_INDICATEURS = "Data_Hackathon.csv"
FICHIER_ETABLISSEMENTS = "annuaire_etablissement.xlsx"
FICHIERS = (FICHIER_DESSERTE, FICHIER_INDICATEURS)

# Jours régénérés à partir d'une semaine type (inclus). Couvre la date de démonstration 2026-11-12.
FENETRE_SEMAINE_TYPE = (date(2026, 11, 1), date(2026, 11, 30))

MOIS = ["janvier", "fevrier", "mars", "avril", "mai", "juin", "juillet", "aout",
        "septembre", "octobre", "novembre", "decembre"]


def trouver_dossier() -> Path | None:
    for d in config.dossiers_donnees():
        if all((d / f).exists() for f in FICHIERS):
            return d
    return None


def fichiers_disponibles() -> bool:
    return trouver_dossier() is not None


# --- utilitaires -------------------------------------------------------------------------------

def _pct(v) -> float | None:
    """'63,20%' -> 63.2 (le signe % indique déjà l'échelle 0-100) ; '' -> None."""
    if v is None:
        return None
    s = str(v).strip().replace("%", "").replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _nb(v) -> float | None:
    return _pct(v)


def _heure(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, (time, datetime)):
        return f"{v.hour:02d}:{v.minute:02d}"
    m = re.match(r"\s*(\d{1,2}):(\d{2})", str(v))
    return f"{int(m.group(1)) % 24:02d}:{m.group(2)}" if m else None


def _train(v) -> str:
    s = str(v).strip().replace(" ", "")
    return s.split(".")[0] if re.fullmatch(r"\d+\.0+", s) else s


def _uic(code: str) -> str | None:
    m = re.search(r"(\d{8})$", code or "")
    return m.group(1) if m else None


# --- coordonnées des gares -----------------------------------------------------------------------

def _coordonnees_gtfs(racine: Path) -> dict[str, tuple[float, float]]:
    """nom de commune normalisé -> (lat, lon) du meilleur arrêt « gare » du réseau urbain."""
    fichier = racine / "reseau_urbain_gare_coordonnees" / "stops.txt"
    if not fichier.exists():
        return {}
    meilleur: dict[str, tuple[int, float, float]] = {}
    centre: dict[str, list[float]] = {}
    with open(fichier, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            nom = commun.normaliser(r.get("stop_name") or "")
            tete = nom.split(" - ")[0].strip()
            try:
                c = centre.setdefault(tete, [0.0, 0.0, 0])
                c[0] += float(r["stop_lat"]); c[1] += float(r["stop_lon"]); c[2] += 1
            except (KeyError, ValueError):
                pass
            if "gare" not in nom:
                continue
            rang = 0 if "sncf" in nom else 1
            try:
                lat, lon = float(r["stop_lat"]), float(r["stop_lon"])
            except (KeyError, ValueError):
                continue
            if tete not in meilleur or rang < meilleur[tete][0]:
                meilleur[tete] = (rang, lat, lon)
    # Repli : centre des arrêts de la commune (approximatif) quand aucun arrêt « gare » n'existe.
    out = {k: (v[0] / v[2], v[1] / v[2]) for k, v in centre.items() if v[2]}
    out.update({k: (v[1], v[2]) for k, v in meilleur.items()})
    return out


def _coordonnees_demo() -> dict[str, tuple[float, float]]:
    return {commun.normaliser(nom): (lat, lon) for _, nom, lat, lon, _ in reseau_demo.GARES.values()}


def _coords_gare(nom: str, gtfs: dict, demo: dict) -> tuple[float | None, float | None, str]:
    n = commun.normaliser(nom)
    # Le réseau de démo (coordonnées de la carte du projet) prime pour les gares qu'il connaît ;
    # le GTFS complète les autres (nom de gare = début du nom de l'arrêt).
    if n in demo:
        return (*demo[n], "demo")
    if n in gtfs:
        return (*gtfs[n], "gtfs")
    tete = n.split(" - ")[0]
    for k, v in gtfs.items():
        if k == tete or k.startswith(tete + " ") or tete.startswith(k + " "):
            return (*v, "gtfs")
    return None, None, ""


# --- lecture des sources ---------------------------------------------------------------------------

def lire_desserte(chemin: Path):
    """Flux (date, train, ordre, uic, nom, arrivee, depart) lu en openpyxl read_only."""
    import openpyxl  # absent de requirements.txt de base : voir requirements-etl

    ws = openpyxl.load_workbook(chemin, read_only=True, data_only=True).active
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not r or r[0] is None or r[1] is None:
            continue
        uic = _uic(str(r[5] or ""))
        if not uic:
            continue
        d = r[0].date() if isinstance(r[0], datetime) else date.fromisoformat(str(r[0])[:10])
        yield (d, _train(r[1]), int(float(r[6] or 0)) + 1, uic, str(r[4] or "").strip(),
               _heure(r[2]), _heure(r[3]))


def lire_indicateurs(chemin: Path):
    with open(chemin, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            cle = {commun.normaliser(k): v for k, v in r.items() if k}
            train = _train(cle.get("n° train", ""))
            try:
                m = MOIS.index(commun.normaliser(cle.get("mois", ""))) + 1
                mois = f"{int(cle['annee'])}-{m:02d}"
            except (ValueError, KeyError):
                continue
            if not train:
                continue
            yield train, mois, cle


def lire_etablissements(chemin: Path):
    import openpyxl

    ws = openpyxl.load_workbook(chemin, read_only=True, data_only=True).active
    it = ws.iter_rows(values_only=True)
    entetes = [str(h) for h in next(it)]
    ix = {h: i for i, h in enumerate(entetes)}
    for r in it:
        if r[ix["Libellé de la région"]] != "Occitanie" or r[ix["Libellé de l'état de l'établissement"]] != "OUVERT":
            continue
        nature = str(r[ix["Libellé de la nature de l'UAI"]] or "")
        n = commun.normaliser(nature)
        if "college" in n:
            type_ = "collège"
        elif "lycee" in n:
            type_ = "lycée"
        elif "ecole" in n:
            type_ = "école"
        else:
            continue
        lat, lon = r[ix["Latitude WGS84"]], r[ix["Longitude WGS84"]]
        if lat is None or lon is None:
            continue
        yield (str(r[ix["Numéro d'UAI"]]), str(r[ix["Appellation officielle"]] or ""), type_,
               str(r[ix["Libellé de la commune"]] or ""), float(lat), float(lon), nature)


# --- semaine type ------------------------------------------------------------------------------------

def jour_type_pour(jours_disponibles: set[date], cible: date) -> date | None:
    """Jour de la même position dans la semaine le plus récent de la desserte (None si aucun)."""
    meme = [d for d in jours_disponibles if d.isoweekday() == cible.isoweekday()]
    return max(meme) if meme else None


def hypotheses(conn) -> list[str]:
    """Hypothèses actives de la base SNCF, pour /api/sante."""
    out = []
    deb, fin = db.meta_get(conn, "semaine_type_du"), db.meta_get(conn, "semaine_type_au")
    if deb and fin:
        out.append(f"Horaires du {deb} au {fin} reconstitués à partir d'une semaine type : pour chaque "
                   "jour de la semaine, circulations du même jour le plus récent de la desserte "
                   "réelle (février à juin 2026).")
    if db.meta_get(conn, "groupes_hypothese") == "1":
        out.append("Aucune liste officielle des trains autorisés aux groupes : tous les trains de la "
                   "desserte sont considérés autorisés ; ceux absents de Data_Hackathon.csv n'ont pas "
                   "d'historique de fiabilité (notes neutres).")
    return out


def _materialiser_semaine_type(conn, jours: set[date]) -> int:
    deb, fin = FENETRE_SEMAINE_TYPE
    n = 0
    d = deb
    while d <= fin:
        ref = jour_type_pour(jours, d)
        if ref and d not in jours:
            cur = conn.execute(
                """INSERT INTO desserte (train_numero, date_circulation, jour_semaine, ordre,
                     gare_uic, heure_arrivee, heure_depart, est_simule)
                   SELECT train_numero, ?, ?, ordre, gare_uic, heure_arrivee, heure_depart, 1
                   FROM desserte WHERE date_circulation = ?""",
                (d.isoformat(), d.isoweekday(), ref.isoformat()))
            n += cur.rowcount
        d += timedelta(days=1)
    return n


# --- construction ---------------------------------------------------------------------------------------

def construire(path: Path | None = None) -> None:
    dossier = trouver_dossier()
    if dossier is None:
        raise FileNotFoundError("Fichiers SNCF introuvables : " + ", ".join(FICHIERS))
    racine = config.REPO_DIR
    path = path or config.db_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        path.unlink()

    print("Lecture de la desserte (xlsx, ~40 s)...", file=sys.stderr)
    lignes, gares, jours = [], {}, set()
    circ_mois: dict[tuple[str, str], set[date]] = defaultdict(set)
    for d, train, ordre, uic, nom, arr, dep in lire_desserte(dossier / FICHIER_DESSERTE):
        gares.setdefault(uic, nom)
        jours.add(d)
        circ_mois[(train, d.strftime("%Y-%m"))].add(d)
        lignes.append((train, d.isoformat(), d.isoweekday(), ordre, uic, arr, dep))

    indicateurs, infos = [], {}
    for train, mois, c in lire_indicateurs(dossier / FICHIER_INDICATEURS):
        infos[train] = ((c.get("axe") or "").strip(), (c.get("origine destination") or "").strip())
        circ = _nb(c.get("circ prevues theorique"))
        nsupp = _nb(c.get("nombre de suppressions")) or 0
        tsupp = _pct(c.get("taux supp"))
        if tsupp is None and circ:  # case vide = aucune suppression
            tsupp = 100.0 * nsupp / circ
        indicateurs.append((
            train, mois, int(circ) if circ is not None else None,
            int(_nb(c.get("circ prevues realise")) or 0), _nb(c.get("frequentation")),
            int(_nb(c.get("nombre de retards pfp")) or 0), _pct(c.get("taux regularite du train")),
            int(nsupp), tsupp,
            _pct(c.get("taux de respect de l emport")), _nb(c.get("minutes perdues"))))

    gtfs, demo = _coordonnees_gtfs(racine), _coordonnees_demo()
    with db.connexion(path) as conn:
        db.creer_schema(conn)
        sans_coord = 0
        for uic, nom in gares.items():
            lat, lon, _ = _coords_gare(nom, gtfs, demo)
            sans_coord += lat is None
            conn.execute("INSERT INTO gare (code_uic, nom, lat, lon, commune, est_simule) "
                         "VALUES (?,?,?,?,NULL,0)", (uic, nom, lat, lon))
        for t in sorted({l[0] for l in lignes} | set(infos)):
            axe, od = infos.get(t, ("", ""))
            orig, _, dest = od.partition("/")
            conn.execute(
                """INSERT INTO train (numero, axe, origine, destination, autorise_groupe,
                     capacite_groupe, est_simule) VALUES (?,?,?,?,?,?,0)""",
                (t, axe, orig.strip(), dest.strip(), 1, config.CAPACITE_GROUPE))
        conn.executemany(
            """INSERT INTO desserte (train_numero, date_circulation, jour_semaine, ordre, gare_uic,
                 heure_arrivee, heure_depart, est_simule) VALUES (?,?,?,?,?,?,?,0)""", lignes)
        conn.executemany(
            """INSERT OR REPLACE INTO indicateur_mensuel (train_numero, mois,
                 circulations_theoriques, circulations_realisees, frequentation, nb_retards,
                 taux_regularite, nb_suppressions, taux_suppression, taux_respect_emport,
                 minutes_perdues, est_simule) VALUES (?,?,?,?,?,?,?,?,?,?,?,0)""",
            indicateurs)
        n_type = _materialiser_semaine_type(conn, jours)
        commun.calculer_fiabilite(conn)
        ignores = commun.inserer_sites(conn, est_simule=0)
        if ignores:
            print("Sites du catalogue sans gare correspondante :", ignores, file=sys.stderr)
        _option_scenario(conn)
        commun.inserer_retours_simules(conn)
        n_etab = _importer_etablissements(conn, dossier)
        import_datatourisme.importer_si_disponible(conn)
        db.meta_set(conn, "donnees", "sncf")
        db.meta_set(conn, "source_libelle", "SNCF 2026")
        db.meta_set(conn, "construite_le", db.maintenant_iso())
        db.meta_set(conn, "sequence_dossier", "0")
        db.meta_set(conn, "source", f"fichiers SNCF ({dossier})")
        db.meta_set(conn, "groupes_hypothese", "1")
        db.meta_set(conn, "semaine_type_du", FENETRE_SEMAINE_TYPE[0].isoformat())
        db.meta_set(conn, "semaine_type_au", FENETRE_SEMAINE_TYPE[1].isoformat())
    print(f"Base SNCF créée : {path} ({len(lignes)} arrêts réels + {n_type} de semaine type, "
          f"{len(gares)} gares dont {sans_coord} sans coordonnées, {len(indicateurs)} indicateurs, "
          f"{n_etab} établissements)", file=sys.stderr)


def _importer_etablissements(conn, dossier: Path) -> int:
    conn.execute("""CREATE TABLE IF NOT EXISTS etablissement (
        uai TEXT PRIMARY KEY, nom TEXT NOT NULL, type TEXT, commune TEXT, lat REAL, lon REAL,
        nature TEXT)""")
    f = dossier / FICHIER_ETABLISSEMENTS
    if not f.exists():
        return 0
    try:
        conn.executemany("INSERT OR REPLACE INTO etablissement VALUES (?,?,?,?,?,?,?)",
                         list(lire_etablissements(f)))
    except ImportError:
        print("openpyxl absent : établissements ignorés.", file=sys.stderr)
        return 0
    return conn.execute("SELECT COUNT(*) FROM etablissement").fetchone()[0]


def _option_scenario(conn) -> None:
    """Option simulée d'un autre groupe sur un train du matin Foix -> Toulouse le 12/11/2026."""
    g = {commun.normaliser(r["nom"]): r["code_uic"] for r in conn.execute("SELECT * FROM gare")}
    foix, toulouse = g.get("foix"), g.get("toulouse matabiau")
    if not foix or not toulouse:
        return
    row = conn.execute(
        """SELECT a.train_numero FROM desserte a
           JOIN desserte b ON b.train_numero = a.train_numero
                AND b.date_circulation = a.date_circulation AND b.ordre > a.ordre
           JOIN train t ON t.numero = a.train_numero AND t.autorise_groupe = 1
           WHERE a.gare_uic = ? AND b.gare_uic = ? AND a.date_circulation = '2026-11-12'
             AND a.heure_depart >= '08:00' ORDER BY a.heure_depart LIMIT 1""",
        (foix, toulouse)).fetchone()
    if row:
        conn.execute("""INSERT INTO positionnement (train_numero, date_sortie, nb_places,
                          dossier_reference, est_simule) VALUES (?, '2026-11-12', 10, NULL, 1)""",
                     (row["train_numero"],))


if __name__ == "__main__":
    construire()
