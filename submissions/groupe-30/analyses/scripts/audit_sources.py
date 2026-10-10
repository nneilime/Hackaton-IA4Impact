"""Audit reproductible des sources et du notebook. Aucun modèle final ni soumission."""
from pathlib import Path
from decimal import Decimal, InvalidOperation
import hashlib
import json
import math
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "analyses"
KIT = ROOT / "DATA 2 _le kit étudiants" / "BatIA_Kit_Hackathon"
D = {p.stem: pd.read_csv(p, dtype=str, keep_default_na=False, encoding="utf-8-sig")
     for p in sorted((KIT / "donnees").glob("*.csv"))}
r, b, l, e, d, i, c, prest = [D[x] for x in ["01_residences", "02_batiments", "03_logements",
    "04_equipements", "05_demandes_locataires", "06_interventions", "07_commandes", "09_prestataires"]]
q = e[e.dans_perimetre_prediction.eq("Oui")].copy()
it = pd.to_datetime(i.date_intervention)
dc = pd.to_datetime(d.date_creation)
cl = pd.to_datetime(d.date_cloture, errors="coerce")
rep = i[i.type_intervention.eq("Remplacement")].copy()
rep["date"] = pd.to_datetime(rep.date_intervention)
corrective = i.type_intervention.eq("Dépannage")
blocking = corrective & i.gravite.eq("Arrêt total")

result = {"scope": {"equipments": len(e), "to_score": len(q),
    "families": q.famille.value_counts().to_dict(), "residences": q.residence_id.value_counts().to_dict(),
    "levels": q.niveau_rattachement.value_counts().to_dict()}}

keys = {"01_residences":"residence_id", "02_batiments":"batiment_id", "03_logements":"logement_id",
        "04_equipements":"equipement_id", "05_demandes_locataires":"demande_id", "06_interventions":"intervention_id",
        "07_commandes":"commande_id", "08_plan_investissement":"operation_id", "09_prestataires":"prestataire_id",
        "10_ref_codes_metiers":"code_metier", "10_ref_familles_equipements":"famille", "10_ref_seuils_dpe":"classe"}
result["primary_keys"] = {name:{"duplicates":int(D[name][col].duplicated().sum()),
    "empty":int(D[name][col].eq("").sum())} for name,col in keys.items()}
result["motifs_composite_key_duplicates"] = int(D["10_ref_motifs_demandes"].duplicated(["motif","parties"]).sum())

ref = {"residence_id":set(r.residence_id), "batiment_id":set(b.batiment_id), "logement_id":set(l.logement_id),
       "equipement_id":set(e.equipement_id), "prestataire_id":set(prest.prestataire_id), "demande_id_liee":set(d.demande_id),
       "code_metier":set(D["10_ref_codes_metiers"].code_metier), "famille":set(D["10_ref_familles_equipements"].famille)}
result["foreign_keys"] = {name:{col:int((df[col].ne("") & ~df[col].isin(values)).sum())
    for col,values in ref.items() if col in df and col != keys.get(name)} for name,df in D.items()}

def mismatch(df, key, parent, parent_key, columns):
    joined = df[df[key].ne("")].merge(parent[[parent_key]+columns], left_on=key,
        right_on=parent_key, how="left", suffixes=("", "_parent"), validate="many_to_one")
    return {col:int(joined[col].ne(joined[col+"_parent"]).sum()) for col in columns}

result["location_consistency"] = {
    "logements_batiments":mismatch(l,"batiment_id",b,"batiment_id",["residence_id"]),
    "equipements_logements":mismatch(e,"logement_id",l,"logement_id",["residence_id","batiment_id"]),
    "interventions_equipements":mismatch(i,"equipement_id",e,"equipement_id",["residence_id","batiment_id","logement_id","code_metier"]),
    "demandes_logements":mismatch(d,"logement_id",l,"logement_id",["residence_id","batiment_id"])}

def normalize(value):
    if value is None or value == "":
        return ""
    text = str(value)
    try:
        return str(Decimal(text).normalize())
    except InvalidOperation:
        return text

wb = load_workbook(KIT / "BatIA_Hackathon_Donnees_Fictives.xlsx", read_only=True, data_only=False)
comparisons = {}
for name, df in D.items():
    sheetname = name.replace("10_", "") if name.startswith("10_") else name
    rows = wb[sheetname].iter_rows(values_only=True)
    headers = list(next(rows))
    mismatch_count = 0
    row_count = 0
    for raw, expected in zip(rows, df.itertuples(index=False, name=None)):
        row_count += 1
        mismatch_count += sum(normalize(a) != normalize(b) for a,b in zip(raw,expected))
    comparisons[name] = {"headers_identical":headers==list(df.columns),
        "rows_identical":wb[sheetname].max_row-1==len(df)==row_count, "different_cells":mismatch_count}
wb.close()
result["xlsx_csv_comparison"] = comparisons

housing_candidates = q.groupby("logement_id").size()
result["requests"] = {"communal_requests":int(d.parties.eq("Communes").sum()),
    "communal_requests_with_housing":int((d.parties.eq("Communes") & d.logement_id.ne("")).sum()),
    "linked_corrective":int((corrective & i.demande_id_liee.ne("")).sum()), "all_corrective":int(corrective.sum()),
    "distinct_linked_requests":int(i.loc[i.demande_id_liee.ne(""),"demande_id_liee"].nunique()),
    "housing_with_multiple_scored_equipment":int((housing_candidates[housing_candidates.index != ""] > 1).sum()),
    "technical_texts":int(d.texte_demande.nunique()), "technician_texts_nonempty":int(i.loc[i.commentaire_technicien.ne(""),"commentaire_technicien"].nunique())}

result["replacement"] = {"events":len(rep), "affected_ids":int(rep.equipement_id.nunique()),
    "multiple_replacements_ids":int(rep.equipement_id.value_counts().gt(1).sum()),
    "multiple_replacements_scored_ids":int(rep[rep.equipement_id.isin(q.equipement_id)].equipement_id.value_counts().gt(1).sum()),
    "current_pose_differs_from_last_replacement_day":int((pd.to_datetime(e.set_index("equipement_id").date_pose).reindex(rep.groupby("equipement_id").date.max().index)
        != rep.groupby("equipement_id").date.max().dt.normalize()).sum())}

result["date_quality"] = {"dpe_after_2024":int(pd.to_datetime(l.dpe_date).gt("2024-12-31").sum()),
    "occupant_entry_non_date_values":l.loc[pd.to_datetime(l.date_entree_locataire_actuel,errors="coerce",format="mixed").isna(),"date_entree_locataire_actuel"].value_counts().to_dict(),
    "request_closure_before_creation":int(cl.lt(dc).sum()),
    "interventions_after_2025":int(it.ge("2026-01-01").sum()),
    "requests_after_2025":int(dc.ge("2026-01-01").sum())}

# Reproduire les variables fournies, afin d'en mesurer les limites sans les corriger silencieusement.
def notebook_variables(date_ref):
    cutoff = pd.Timestamp(date_ref)
    h = i[it.le(cutoff)]
    X = q[["equipement_id","famille","duree_vie_reference_ans"]].copy()
    rempl = pd.to_datetime(h[h.type_intervention.eq("Remplacement")].set_index("equipement_id").date_intervention)
    rempl = rempl.groupby(level=0).max()
    pose = pd.to_datetime(q.set_index("equipement_id").date_pose)
    prev = pd.to_datetime(q.set_index("equipement_id").date_pose_precedente,errors="coerce")
    pose_ref = pd.Series(pd.to_datetime(np.where(pose.le(cutoff),pose,prev)),index=pose.index).fillna(pose)
    pose_ref = pd.concat([pose_ref,rempl],axis=1).max(axis=1)
    X["age_ans"] = (cutoff-X.equipement_id.map(pose_ref)).dt.days/365.25
    X["age_relatif"] = X.age_ans/X.duree_vie_reference_ans.astype(float)
    dep = h[h.type_intervention.eq("Dépannage") & pd.to_datetime(h.date_intervention).gt(cutoff-pd.Timedelta(days=365))]
    X["nb_depannages_12m"] = X.equipement_id.map(dep.groupby("equipement_id").size()).fillna(0)
    return X,pose_ref

def target(year, complete_last_day):
    end = pd.Timestamp(f"{year}-10-01") if complete_last_day else pd.Timestamp(f"{year}-09-30")
    mask = blocking & it.ge(f"{year}-01-01") & (it.lt(end) if complete_last_day else it.le(end))
    return q.equipement_id.isin(i.loc[mask,"equipement_id"]).astype(int).to_numpy()

cutoff_audits = {}
for year in range(2017,2026):
    cutoff = pd.Timestamp(f"{year}-12-31")
    X,pose_ref = notebook_variables(cutoff)
    prior_replacements = rep[rep.date.lt(cutoff+pd.Timedelta(days=1))]
    total_reps_at_cutoff = q.equipement_id.map(prior_replacements.groupby("equipement_id").size()).fillna(0).astype(int)
    last_rep = prior_replacements.groupby("equipement_id").date.max()
    dep_last_year = i[corrective & it.gt(cutoff-pd.Timedelta(days=365)) & it.le(cutoff)]
    old_lifetime = (pd.to_datetime(dep_last_year.date_intervention).lt(dep_last_year.equipement_id.map(last_rep))
        & dep_last_year.equipement_id.isin(q.equipement_id)) if len(last_rep) else pd.Series(False,index=dep_last_year.index)
    cutoff_audits[str(year)] = {"negative_age_ids":X.loc[X.age_ans.lt(0),"equipement_id"].tolist(),
        "events_excluded_on_dec31":int((it.ge(cutoff)&it.lt(cutoff+pd.Timedelta(days=1))).sum()),
        "current_replacement_count_leaks_future_ids":int(q.nb_remplacements_2018_2025.astype(int).ne(total_reps_at_cutoff).sum()),
        "requests_open_at_cutoff_with_future_closure":int((dc.lt(cutoff+pd.Timedelta(days=1))&cl.ge(cutoff+pd.Timedelta(days=1))).sum()),
        "old_lifetime_corrective_events_in_12m_feature":int(old_lifetime.sum()),
        "old_lifetime_equipment_ids":int(dep_last_year.loc[old_lifetime,"equipement_id"].nunique())}
result["notebook_cutoff_audits"] = cutoff_audits

labels = {}
for year in range(2018,2026):
    all_events = blocking & it.ge(f"{year}-01-01") & it.lt(f"{year}-10-01") & i.equipement_id.isin(q.equipement_id)
    counts = i.loc[all_events].groupby("equipement_id").size()
    labels[str(year)] = {"blocking_events":int(all_events.sum()),"equipment_with_blocking_event":int(len(counts)),
        "prevalence":len(counts)/len(q), "notebook_target_omits_ids":int((target(year,True)>target(year,False)).sum())}
result["target_history"] = labels

X24,_ = notebook_variables("2024-12-31")
y25 = target(2025,False)
def metrics(scores, y):
    k = math.ceil(len(scores)*.1)
    order = np.lexsort((q.equipement_id.to_numpy(),-np.asarray(scores)))
    hits = int(y[order[:k]].sum())
    positive = np.asarray(scores)[y==1]
    negative = np.asarray(scores)[y==0]
    # Définition de l'AUC par comparaisons de paires, avec demi-crédit aux ex aequo.
    auc = ((positive[:,None]>negative).sum()+.5*(positive[:,None]==negative).sum())/(len(positive)*len(negative))
    return {"auc":float(auc), "top_k":k,"hits":hits,
        "precision_top10":hits/k,"recall_top10":hits/int(y.sum()),"lift_top10":(hits/k)/y.mean()}
result["notebook_baselines_reproduced"] = {"target_positives":int(y25.sum()),
    "age_relative":metrics(X24.age_relatif,y25),
    "logistic_2_features":"Non exécutée : scikit-learn absent du runtime. Aucun paquet installé pour cet audit de cadrage.",
    "warning":"Reproduction du notebook avec ses limites temporelles ; aucune performance finale ni preuve terrain."}
# Comparaison indicative après inclusion des journées entières, sans entraînement.
_,pose25 = notebook_variables("2024-12-31 23:59:59.999999999")
correct_age = (pd.Timestamp("2024-12-31") - q.equipement_id.map(pose25).dt.normalize()).dt.days/365.25
result["age_baseline_full_days"] = metrics(correct_age/q.duree_vie_reference_ans.astype(float),target(2025,True))

family_events = i[blocking & it.ge("2025-01-01") & it.lt("2025-10-01")]
family_labels = q.assign(y=q.equipement_id.isin(family_events.equipement_id).astype(int)).groupby("famille").agg(n=("equipement_id","size"),positive=("y","sum"))
result["target_2025_by_family"] = family_labels.to_dict(orient="index")

quarterly = {}
for quarter,start,end in [(2,"2025-04-01","2025-07-01"),(3,"2025-07-01","2025-10-01"),(4,"2025-10-01","2026-01-01")]:
    h = i[corrective & it.ge(start) & it.lt(end)]
    # MUPPR couvre également des appareils hors contrat dans le CSV.
    # Les UO du rapport correspondent au sous-ensemble couvert par le contrat.
    h = h[h.code_metier.isin(["CHGAZ","CHPAC","CHFRC","CHVBO","CHVCC","MUPPR"])
          & h.equipement_id.isin(e.loc[e.sous_contrat_entretien.eq("Oui"),"equipement_id"])]
    texts = (OUT/"sources_lisibles"/f"Rapport_activite_T{quarter}_2025_FICTIF.md").read_text(encoding="utf-8")
    observations = texts.split("Observations du technicien référent")[-1]
    ids = re.findall(r"EQP-\d+",observations)
    end_day = pd.Timestamp(end)
    replaced = rep[rep.date.lt(end_day) & rep.equipement_id.isin(ids)]
    quarterly[str(quarter)] = {"corrective_events":len(h),"blocking_events":int(h.gravite.eq("Arrêt total").sum()),
        "by_trade":h.groupby("code_metier").size().to_dict(), "alert_equipment_ids":ids,
        "alert_ids_outside_prediction_scope":sorted(set(ids)-set(q.equipement_id)),
        "alert_ids_with_replacement_in_quarter":sorted(set(replaced.loc[replaced.date.ge(start),"equipement_id"])),
        "alert_ids_with_replacement_in_prior_9_months":sorted(set(replaced.loc[replaced.date.ge(end_day-pd.DateOffset(months=9)),"equipement_id"])),
        "codes_absent_from_report":["ASC","CHGCO","PECSO","MUPCO"]}
result["report_cross_checks"] = quarterly

manifest = json.loads((OUT/"inventaire_sources.json").read_text(encoding="utf-8"))
result["source_hashes_unchanged"] = all(hashlib.sha256((ROOT/x["path"]).read_bytes()).hexdigest()==x["sha256"] for x in manifest)
assert result["scope"]["to_score"] == 1285, "Le périmètre ne correspond plus au contrat du kit."
assert result["source_hashes_unchanged"], "Une source a changé depuis l'extraction."
assert all(x["headers_identical"] and x["rows_identical"] and x["different_cells"]==0 for x in comparisons.values()), "Les CSV et Excel divergent."
assert all(x["duplicates"]==x["empty"]==0 for x in result["primary_keys"].values()), "Clé primaire invalide."
assert all(count==0 for checks in result["foreign_keys"].values() for count in checks.values()), "Référence absente."
assert [(quarterly[str(k)]["corrective_events"],quarterly[str(k)]["blocking_events"]) for k in [2,3,4]] == [(225,39),(223,32),(276,25)], "Écart avec les tableaux des rapports PDF."
rank_auc = (pd.Series(X24.age_relatif.to_numpy()).rank().to_numpy()[y25==1].sum()-y25.sum()*(y25.sum()+1)/2)/(y25.sum()*(len(y25)-y25.sum()))
assert abs(rank_auc-result["notebook_baselines_reproduced"]["age_relative"]["auc"]) < 1e-12, "Calculs indépendants de l'AUC en désaccord."
result["validation"] = {"integrity_and_scope":"PASS", "auc_pairwise_vs_ranks":"PASS",
                        "pdf_quarterly_totals":"PASS", "original_source_hashes":"PASS"}
(OUT/"audit_resultats.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")

profile = json.loads((OUT/"profil_donnees.json").read_text(encoding="utf-8"))
lines = ["# Audit des données BatIA", "", "Périmètre : prévention des pannes. Résultats calculés localement par `scripts/audit_sources.py`. Les anomalies sont documentées sans modification des sources.",
         "", "## Intégrité et couverture", "", "| Table | Lignes | Colonnes |", "| --- | ---: | ---: |"]
for name,v in profile.items():
    lines.append(f"| {name} | {v['rows']:,} | {len(v['columns'])} |".replace(","," "))
lines.extend(["", "- Les clés primaires vérifiées sont uniques et renseignées. Aucun identifiant non vide ne pointe vers une référence absente dans les contrôles réalisés.",
    "- Les localisations redondantes vérifiées entre logements, bâtiments, équipements, interventions et demandes concordent.",
    "- Les 13 CSV ont les mêmes en-têtes, nombres de lignes et valeurs que les onglets Excel correspondants après normalisation numérique.",
    "- La clé du référentiel des motifs est le couple `(motif, parties)`. Le motif VMC apparaît pour les parties privatives et communes.",
    "- Les 24 empreintes SHA-256 des sources sont inchangées. Le fichier temporaire Excel préexistant est exclu.",
    "", "## Pannes observées dans notre périmètre", "", "| Janvier-septembre | Arrêts totaux (événements) | Équipements positifs | Prévalence | Positifs omis par la borne du notebook |",
    "| --- | ---: | ---: | ---: | ---: |"])
for year,v in labels.items():
    lines.append(f"| {year} | {v['blocking_events']} | {v['equipment_with_blocking_event']} | {v['prevalence']:.1%} | {v['notebook_target_omits_ids']} |")
lines.extend(["", "La cible est binaire par équipement. Les événements répétés ne sont pas des équipements supplémentaires. Les très petits effectifs par famille limitent l'interprétation.",
    "", "| Famille | Équipements | Positifs janvier-septembre 2025 |", "| --- | ---: | ---: |"])
for family,v in result["target_2025_by_family"].items():
    lines.append(f"| {family} | {v['n']} | {v['positive']} |")
lines.extend(["", "## Audit temporel du notebook", "", "| Référence au 31/12 | Âges négatifs | Passages du dernier jour omis | Appareils dont le cumul actuel de remplacements inclut le futur | Dépannages de l'ancien appareil dans le compteur 12 mois |",
    "| --- | ---: | ---: | ---: | ---: |"])
for year,v in cutoff_audits.items():
    lines.append(f"| {year} | {len(v['negative_age_ids'])} | {v['events_excluded_on_dec31']} | {v['current_replacement_count_leaks_future_ids']} | {v['old_lifetime_corrective_events_in_12m_feature']} |")
lines.extend(["", "L'année 2017 sert uniquement à vérifier la reconstruction avant le début des interventions. Les âges négatifs des validations anciennes montrent que la pose immédiatement précédente ne suffit pas à retrouver tous les premiers appareils.",
    "", "Le notebook ne se sert pas du cumul actuel de remplacements dans ses deux variables : le risque de fuite apparaît si l'on ajoute directement ce champ à un modèle. Son compteur de dépannages mélange en revanche effectivement les cycles de vie.",
    "", "135 DPE ont été établis après fin 2024. Les 347 valeurs `Avant 2018` des entrées d'occupants sont des valeurs textuelles prévues par le dictionnaire, pas des dates corrompues. Une demande ouverte à fin 2023 possède une clôture future. Les réouvertures et échanges ne sont pas horodatés individuellement.",
    "", "## Demandes et textes", "", f"- {result['requests']['linked_corrective']} dépannages sur {result['requests']['all_corrective']} portent une demande liée ({result['requests']['linked_corrective']/result['requests']['all_corrective']:.1%}), pour 362 demandes distinctes.",
    "- 370 logements ont plusieurs appareils du périmètre de prédiction. Une jointure par logement ne suffit pas à attribuer un symptôme.",
    "- Les 968 demandes indiquées comme communes possèdent toutes un logement renseigné. Leur portée doit aussi utiliser `parties` et le motif.",
    f"- {result['requests']['technician_texts_nonempty']} formulations non vides distinctes dans les commentaires des techniciens. La diversité est celle de modèles de simulation.",
    "", "## Remplacements et rapports PDF", "", "866 remplacements concernent 846 identifiants. 20 identifiants ont plusieurs remplacements, dont 12 dans le périmètre de prédiction. Pour tous les emplacements remplacés, la pose actuelle correspond au jour du dernier remplacement.",
    "", "Les rapports couvrent six codes et les équipements sous contrat. Les appareils hors contrat partageant `MUPPR` doivent être exclus pour réconcilier les tableaux.",
    "", "| Rapport | Dépannages réconciliés | Arrêts totaux | Alertes listées | Alertes hors soumission | Appareils alertés remplacés dans les neuf mois résumés |",
    "| --- | ---: | ---: | ---: | ---: | ---: |"])
for quarter,v in quarterly.items():
    lines.append(f"| T{quarter} 2025 | {v['corrective_events']} | {v['blocking_events']} | {len(v['alert_equipment_ids'])} | {len(v['alert_ids_outside_prediction_scope'])} | {len(v['alert_ids_with_replacement_in_prior_9_months'])} |")
base = result["notebook_baselines_reproduced"]["age_relative"]
full = result["age_baseline_full_days"]
lines.extend(["", "Les totaux correspondent aux PDF. L'absence d'un appareil dans les observations ne constitue pas un constat de bon fonctionnement. La mention dans un rapport ne précise pas à elle seule la date du symptôme ou le cycle de vie concerné.",
    "", "## Baseline d'âge et limites de validation", "", f"Âge relatif seul, au 31/12/2024, cible du notebook en janvier-septembre 2025 : **AUC {base['auc']:.6f}**, **{base['hits']}/{base['top_k']}** équipements positifs dans le top 10 %, précision **{base['precision_top10']:.1%}**, rappel **{base['recall_top10']:.1%}**, lift **{base['lift_top10']:.2f}**.",
    "", f"Après inclusion des journées complètes et normalisation des dates de pose : AUC **{full['auc']:.6f}**, mêmes 16 positifs dans les 129 premiers appareils. Ce contrôle ne corrige pas toute la reconstruction des validations anciennes.",
    "", "Convention de cet audit : top 10 % arrondi au supérieur, soit 129 appareils, et ex aequo départagés par identifiant. Le jury peut employer une autre convention. L'AUC a été vérifiée par deux méthodes indépendantes, comparaisons de paires et rangs.",
    "", "La baseline logistique n'a pas été exécutée, car scikit-learn est absent du runtime. Aucun modèle final, aucune prédiction 2026 et aucun CSV de soumission ne sont produits dans cette phase de cadrage.",
    "", "## Sources et reproductibilité", "", "- `audit_resultats.json` : contrôles et résultats détaillés, dont les identifiants concernés.",
    "- `profil_donnees.json` : colonnes, catégories, valeurs manquantes, exemples et bornes de dates.",
    "- `profil_classeur.json` : contenu documentaire et structure du classeur.",
    "- `inventaire_sources.json` : fichiers et empreintes SHA-256.",
    "- [Consignes extraites](sources_lisibles/LISEZ-MOI_BatIA.md), [dictionnaire](sources_lisibles/xlsx_DICTIONNAIRE.csv), [hypothèses](sources_lisibles/xlsx_HYPOTHESES.csv) et [notebook](sources_lisibles/BatIA_notebook_demarrage.md).",
    "- Rapports [T2](sources_lisibles/Rapport_activite_T2_2025_FICTIF.md), [T3](sources_lisibles/Rapport_activite_T3_2025_FICTIF.md), [T4](sources_lisibles/Rapport_activite_T4_2025_FICTIF.md).",
    "", "Les contrôles prouvent les points décrits ici. Ils ne constituent pas une certification de toutes les valeurs simulées ni une validation de performance sur données réelles."])
(OUT/"AUDIT_DONNEES.md").write_text("\n".join(lines)+"\n",encoding="utf-8")
print(json.dumps(result,ensure_ascii=True,indent=2))
