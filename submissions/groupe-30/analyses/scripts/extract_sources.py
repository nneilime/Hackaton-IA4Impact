"""Extraction non destructive du kit BatIA (Python fourni par Codex)."""
from pathlib import Path
import csv
import hashlib
import json
import re
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd
from pypdf import PdfReader
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "DATA 2 _le kit étudiants"
OUT = ROOT / "analyses"
TEXT = OUT / "sources_lisibles"
TEXT.mkdir(parents=True, exist_ok=True)

def write_json(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

manifest = []
for file in sorted(SOURCE.rglob("*")):
    if not file.is_file() or file.name.startswith("~$"):
        continue
    manifest.append({"path": str(file.relative_to(ROOT)), "bytes": file.stat().st_size,
                     "sha256": hashlib.sha256(file.read_bytes()).hexdigest()})
write_json("inventaire_sources.json", manifest)

for file in sorted(SOURCE.rglob("*.pdf")):
    reader = PdfReader(file)
    parts = [f"# Source : {file.relative_to(ROOT)}", f"Pages : {len(reader.pages)}"]
    for n, page in enumerate(reader.pages, 1):
        extracted = (page.extract_text() or "[Aucun texte extrait : vérifier visuellement]").replace("\x7f", "-")
        parts.extend([f"\n## Page {n}\n", extracted])
    (TEXT / f"{file.stem}.md").write_text("\n".join(parts), encoding="utf-8")

for file in SOURCE.rglob("*.pptx"):
    parts = [f"# Source : {file.relative_to(ROOT)}"]
    with zipfile.ZipFile(file) as archive:
        slides = sorted((x for x in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", x)),
                        key=lambda x: int(re.search(r"slide(\d+)", x).group(1)))
        for slide in slides:
            tree = ET.fromstring(archive.read(slide))
            text = [e.text for e in tree.iter() if e.tag.endswith("}t") and e.text]
            parts.extend([f"\n## {slide}\n", "\n".join(text)])
        for note in sorted(x for x in archive.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", x)):
            tree = ET.fromstring(archive.read(note))
            parts.extend([f"\n## Notes {note}\n", "\n".join(e.text for e in tree.iter() if e.tag.endswith("}t") and e.text)])
    (TEXT / f"{file.stem}_pptx.md").write_text("\n".join(parts), encoding="utf-8")

for file in SOURCE.rglob("*.ipynb"):
    notebook = json.loads(file.read_text(encoding="utf-8"))
    parts = [f"# Source : {file.relative_to(ROOT)}"]
    for i, cell in enumerate(notebook["cells"], 1):
        content = "".join(cell.get("source", []))
        parts.append(f"\n## Cellule {i} ({cell['cell_type']})\n")
        parts.append(f"```python\n{content}\n```" if cell["cell_type"] == "code" else content)
        for output in cell.get("outputs", []):
            if "text" in output:
                parts.append("".join(output["text"]))
            if "text/plain" in output.get("data", {}):
                parts.append("".join(output["data"]["text/plain"]))
    (TEXT / f"{file.stem}.md").write_text("\n".join(parts), encoding="utf-8")

profiles = {}
for file in sorted(SOURCE.rglob("*.csv")):
    raw = file.read_text(encoding="utf-8-sig")
    delimiter = csv.Sniffer().sniff(raw[:10000], delimiters=";,\t").delimiter
    df = pd.read_csv(file, sep=delimiter, encoding="utf-8-sig", dtype=str, keep_default_na=False)
    fields = {}
    for col in df.columns:
        values = df[col]
        notempty = values[values != ""]
        fields[col] = {"empty": int((values == "").sum()), "unique": int(notempty.nunique()),
                       "examples": notempty.drop_duplicates().head(6).tolist()}
        if notempty.nunique() <= 30:
            fields[col]["counts"] = values.value_counts(dropna=False).to_dict()
        if col.startswith("date") or col.endswith("_date"):
            parsed = pd.to_datetime(notempty, errors="coerce", format="mixed")
            fields[col]["date_min"] = str(parsed.min())
            fields[col]["date_max"] = str(parsed.max())
            fields[col]["invalid_dates"] = int(parsed.isna().sum())
    profiles[file.stem] = {"path": str(file.relative_to(ROOT)), "delimiter": delimiter,
                           "rows": len(df), "columns": fields, "exact_duplicate_rows": int(df.duplicated().sum())}
write_json("profil_donnees.json", profiles)

workbooks = {}
for file in SOURCE.rglob("*.xlsx"):
    if file.name.startswith("~$"):
        continue
    wb = load_workbook(file, read_only=True, data_only=False)
    sheets = {}
    for sheet in wb:
        rows = list(sheet.iter_rows(values_only=True))
        sheets[sheet.title] = {"rows": sheet.max_row, "columns": sheet.max_column,
                               "first_rows": rows[:8],
                               "formulas": sum(1 for row in rows for value in row if isinstance(value, str) and value.startswith("="))}
        if len(rows) < 200:
            sheets[sheet.title]["all_rows"] = rows
        # Les tables sont déjà fournies en CSV : n'exporter que la documentation.
        if sheet.title in {"LISEZMOI", "DICTIONNAIRE", "HYPOTHESES"}:
            with (TEXT / f"xlsx_{sheet.title}.csv").open("w", encoding="utf-8", newline="") as stream:
                csv.writer(stream).writerows(rows)
    workbooks[file.name] = sheets
    wb.close()
write_json("profil_classeur.json", workbooks)
print(json.dumps({"sources": len(manifest), "pdf": len(list(SOURCE.rglob('*.pdf'))),
                  "csv": {name: data['rows'] for name, data in profiles.items()},
                  "xlsx": {name: list(sheets) for name, sheets in workbooks.items()}}, ensure_ascii=False, indent=2))
